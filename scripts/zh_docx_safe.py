#!/usr/bin/env python
"""Generate Chinese-friendly DOCX reports from a UTF-8 JSON spec."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except Exception as exc:  # pragma: no cover - environment guard
    raise SystemExit(
        "python-docx is required. Install it or use the bundled environment that provides it."
    ) from exc


DEFAULT_CN_FONT = "微软雅黑"
DEFAULT_EN_FONT = "Arial"

EXAMPLE_SPEC: dict[str, Any] = {
    "title": "中文 DOCX 稳定生成示例",
    "subtitle": "zh-docx-safe example",
    "output": "D:/tmp/CX_中文DOCX稳定生成示例.docx",
    "sections": [
        {"heading": "核心判断", "level": 1},
        {"paragraph": "这是一份包含中文路径、中文正文、中英混排和表格的 Word 生成示例。"},
        {
            "note_box": {
                "title": "注意",
                "items": [
                    "用 UTF-8 JSON 作为输入。",
                    "运行 Python 时使用 -X utf8。",
                    "生成后自动校验 DOCX zip 完整性。",
                ],
            }
        },
        {"quote": "If Chinese text survives the whole path from JSON to DOCX, the workflow is stable."},
        {
            "table": {
                "headers": ["项目", "预期结果"],
                "rows": [
                    ["中文路径", "正常创建输出目录和文件"],
                    ["中文字体", "使用微软雅黑作为默认东亚字体"],
                ],
            }
        },
    ],
    "sources": [["zh-docx-safe", "https://github.com/baihe26/zh-docx-safe"]],
}


def set_east_asian_font(run: Any, font_name: str) -> None:
    run.font.name = DEFAULT_EN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_border(paragraph: Any, color: str = "B7D4E8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def style_doc(doc: Document, cn_font: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DEFAULT_EN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    normal.font.size = Pt(10.5)

    for style_name, size in [
        ("Title", 18),
        ("Heading 1", 15),
        ("Heading 2", 12.5),
        ("Heading 3", 11.5),
    ]:
        style = styles[style_name]
        style.font.name = DEFAULT_EN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
        style.font.size = Pt(size)
        style.font.bold = True


def add_run(paragraph: Any, text: str, cn_font: str, size: float | None = None, bold: bool = False) -> Any:
    run = paragraph.add_run(text)
    set_east_asian_font(run, cn_font)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    return run


def add_title(doc: Document, title: str, subtitle: str | None, cn_font: str) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(paragraph, title, cn_font, size=18, bold=True)
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = add_run(sub, subtitle, cn_font, size=10)
        run.font.color.rgb = RGBColor(90, 90, 90)


def add_table(doc: Document, table_spec: dict[str, Any], cn_font: str) -> None:
    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    if not headers and rows:
        headers = [f"列{i + 1}" for i in range(len(rows[0]))]
    if not headers:
        return

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    font_size = float(table_spec.get("font_size", 8.5))

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, table_spec.get("header_fill", "D9EAF7"))
        cell.text = ""
        add_run(cell.paragraphs[0], str(header), cn_font, size=font_size, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(headers)]):
            cells[index].text = ""
            add_run(cells[index].paragraphs[0], str(value), cn_font, size=font_size)
    doc.add_paragraph()


def add_note_box(doc: Document, note_spec: dict[str, Any], cn_font: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, note_spec.get("fill", "EEF5F7"))
    cell.text = ""

    title = note_spec.get("title")
    if title:
        add_run(cell.paragraphs[0], str(title), cn_font, size=9.5, bold=True)
    for item in note_spec.get("items", []):
        paragraph = cell.add_paragraph(style="List Bullet")
        add_run(paragraph, str(item), cn_font, size=9)
    if note_spec.get("paragraph"):
        paragraph = cell.add_paragraph()
        add_run(paragraph, str(note_spec["paragraph"]), cn_font, size=9)
    doc.add_paragraph()


def add_quote(doc: Document, text: str, cn_font: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    set_paragraph_border(paragraph)
    run = add_run(paragraph, text, cn_font, size=10)
    run.italic = True


def add_section_item(doc: Document, item: dict[str, Any], cn_font: str) -> None:
    if "heading" in item:
        level = int(item.get("level", 1))
        doc.add_heading(str(item["heading"]), level=level)
    elif "paragraph" in item:
        paragraph = doc.add_paragraph()
        add_run(paragraph, str(item["paragraph"]), cn_font)
    elif "bullets" in item:
        for bullet in item["bullets"]:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_run(paragraph, str(bullet), cn_font)
    elif "numbers" in item:
        for number in item["numbers"]:
            paragraph = doc.add_paragraph(style="List Number")
            add_run(paragraph, str(number), cn_font)
    elif "table" in item:
        add_table(doc, item["table"], cn_font)
    elif "note_box" in item:
        add_note_box(doc, item["note_box"], cn_font)
    elif "quote" in item:
        add_quote(doc, str(item["quote"]), cn_font)
    elif item.get("page_break"):
        doc.add_page_break()


def build_doc(spec: dict[str, Any]) -> Path:
    output_raw = spec.get("output")
    if not output_raw:
        raise ValueError("Spec must include an output path.")
    output = Path(output_raw)
    output.parent.mkdir(parents=True, exist_ok=True)

    cn_font = spec.get("cn_font", DEFAULT_CN_FONT)
    doc = Document()
    style_doc(doc, cn_font)

    if spec.get("title"):
        add_title(doc, str(spec["title"]), spec.get("subtitle"), cn_font)

    for item in spec.get("sections", []):
        add_section_item(doc, item, cn_font)

    if spec.get("sources"):
        doc.add_heading("来源", level=1)
        add_table(
            doc,
            {"headers": ["来源", "链接"], "rows": spec["sources"], "font_size": 8.0},
            cn_font,
        )

    doc.save(output)
    with zipfile.ZipFile(output, "r") as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"DOCX zip validation failed at {bad}")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", help="UTF-8 JSON spec path.")
    parser.add_argument("--write-example", help="Write an example UTF-8 JSON spec to this path and exit.")
    args = parser.parse_args()

    if args.write_example:
        example_path = Path(args.write_example)
        example_path.parent.mkdir(parents=True, exist_ok=True)
        example_path.write_text(
            json.dumps(EXAMPLE_SPEC, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(str(example_path))
        return 0

    if not args.spec:
        parser.error("--spec is required unless --write-example is used.")

    spec_path = Path(args.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    output = build_doc(spec)
    print(str(output))
    return 0


if __name__ == "__main__":
    sys.exit(main())
